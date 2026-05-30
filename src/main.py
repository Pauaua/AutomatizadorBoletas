"""
Interfaz principal para el flujo de boletas SII.
- Pestaña 1: Ingreso individual RUT/Clave.
- Pestaña 2: Carga de Excel (RUT/Clave) y procesamiento masivo.
Toda la lógica de automatización está en core/boleta_automator.py.
"""
import sys
import os
import io

# Fix encoding para terminales Windows (cp1252 no soporta emojis)
if sys.stdout and hasattr(sys.stdout, 'buffer'):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
if sys.stderr and hasattr(sys.stderr, 'buffer'):
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

def get_base_path():
    if getattr(sys, "frozen", False):
        return sys._MEIPASS if hasattr(sys, "_MEIPASS") else os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def get_resource_path(*relative_path):
    if getattr(sys, "frozen", False):
        base = sys._MEIPASS if hasattr(sys, "_MEIPASS") else os.path.dirname(sys.executable)
        return os.path.join(base, "src", *relative_path)
    base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, *relative_path)

SCRIPT_DIR = get_base_path()
if getattr(sys, "frozen", False):
    PROJECT_ROOT = os.path.dirname(sys.executable)
    SRC_DIR = os.path.join(SCRIPT_DIR, "src")
else:
    PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
    SRC_DIR = SCRIPT_DIR

sys.path.insert(0, SRC_DIR if getattr(sys, "frozen", False) else SCRIPT_DIR)

CORE_DIR = get_resource_path("core")
if not os.path.exists(CORE_DIR):
    print(f"❌ ERROR: No existe la carpeta 'core' en {CORE_DIR}")
    sys.exit(1)
AUTOMATOR_FILE = os.path.join(CORE_DIR, "boleta_automator.py")
if not os.path.exists(AUTOMATOR_FILE):
    print(f"❌ ERROR: No existe boleta_automator.py en {CORE_DIR}")
    sys.exit(1)

try:
    import selenium
    print(f"✅ Selenium {selenium.__version__}")
except ImportError:
    print("❌ pip install selenium webdriver-manager PyQt5 pandas openpyxl")
    sys.exit(1)
try:
    import PyQt5
    print("✅ PyQt5 instalado")
except ImportError:
    print("❌ pip install PyQt5")
    sys.exit(1)
try:
    from core.boleta_automator import BoletaAutomator
    print("✅ BoletaAutomator importado")
except Exception as e:
    print(f"❌ Error importando BoletaAutomator: {e}")
    import traceback
    traceback.print_exc()
    sys.exit(1)

try:
    from PyQt5.QtWidgets import (
        QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
        QLabel, QLineEdit, QPushButton, QTextEdit, QCheckBox, QGroupBox,
        QMessageBox, QProgressBar, QSplitter, QTabWidget, QTableWidget,
        QTableWidgetItem, QFileDialog, QHeaderView, QComboBox, QScrollArea,
        QSizePolicy,
    )
    from PyQt5.QtCore import Qt, QTimer
    from PyQt5.QtGui import QFont, QPalette, QColor, QPixmap, QIcon
    print("✅ Componentes PyQt5 importados")
except ImportError as e:
    print(f"❌ Error PyQt5: {e}")
    sys.exit(1)


class BoletaGUI(QMainWindow):
    """Ventana principal: pestaña Individual + pestaña Excel (carga y proceso)."""

    def __init__(self):
        super().__init__()
        self.automator = BoletaAutomator()
        self.worker_individual = None
        self.worker_masivo = None
        self.excel_path = None
        self.excel_df = None
        self.excel_cols = None
        self.masivo_detenido = False
        self.masivo_resultados = []   # acumula dict por fila para reporte final
        self.init_ui()

    # Perfiles de vista:
    # (ancho, alto, font_title, logo_h, splitter_orient, splitter_sizes, log_masivo_h)
    VIEW_PROFILES = {
        "PC":     (1000, 720, 17, 72, "H", [370, 610], 180),
        "Tablet": (768,  900, 14, 52, "V", [440, 400], 150),
        "Móvil":  (430,  800, 11, 40, "V", [420, 340], 120),
    }

    def init_ui(self):
        self.setWindowTitle("Automatizador de Boletas Electrónicas")

        # Ícono: .ico primero, logo.png como fallback
        icon_path = get_resource_path("assets", "icon.ico")
        logo_icon_path = get_resource_path("assets", "logo.png")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))
        elif os.path.exists(logo_icon_path):
            self.setWindowIcon(QIcon(logo_icon_path))

        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # ── Barra de selector de vista ──────────────────────────────
        topbar_widget = QWidget()
        topbar_widget.setStyleSheet("background: #c4cfe0; border-bottom: 1px solid #a3b4cb;")
        topbar_layout = QHBoxLayout(topbar_widget)
        topbar_layout.setContentsMargins(10, 5, 10, 5)
        topbar_layout.setSpacing(6)

        view_lbl = QLabel("Vista:")
        view_lbl.setStyleSheet(
            "color: #365ca3; font-weight: bold; font-size: 12px; background: transparent; border: none;"
        )
        topbar_layout.addWidget(view_lbl)

        self.view_combo = QComboBox()
        self.view_combo.addItems(["🖥️  PC", "📱 Tablet", "📲 Móvil"])
        self.view_combo.setFixedWidth(115)
        self.view_combo.setStyleSheet(
            "QComboBox { padding: 3px 8px; border: 1px solid #a3b4cb; border-radius: 4px;"
            " background: white; color: #365ca3; font-weight: bold; font-size: 12px; }"
            "QComboBox::drop-down { border: none; width: 20px; }"
            "QComboBox QAbstractItemView { background: white; color: #365ca3; }"
        )
        self.view_combo.currentIndexChanged.connect(self._on_view_changed)
        topbar_layout.addWidget(self.view_combo)
        topbar_layout.addStretch()
        main_layout.addWidget(topbar_widget)

        # ── Header: widget contenedor real (logo arriba + título abajo) ──
        self.header_widget = QWidget()
        self.header_widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.header_widget.setStyleSheet("background: #d4dce4;")
        header_layout = QVBoxLayout(self.header_widget)
        header_layout.setContentsMargins(12, 10, 12, 8)
        header_layout.setSpacing(5)
        header_layout.setAlignment(Qt.AlignHCenter)

        # Logo
        self._logo_path = None
        self.logo_label = None
        logo_path = get_resource_path("assets", "logo.png")
        if os.path.exists(logo_path):
            self._logo_path = logo_path
            self.logo_label = QLabel()
            pix = QPixmap(logo_path).scaledToHeight(72, Qt.SmoothTransformation)
            if not pix.isNull():
                self.logo_label.setPixmap(pix)
            self.logo_label.setAlignment(Qt.AlignHCenter | Qt.AlignVCenter)
            self.logo_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            self.logo_label.setStyleSheet("background: transparent;")
            header_layout.addWidget(self.logo_label)

        # Título
        self.title_label = QLabel("Automatizador de Boletas Electrónicas")
        self.title_label.setFont(QFont("Segoe UI", 17, QFont.Bold))
        self.title_label.setAlignment(Qt.AlignHCenter | Qt.AlignVCenter)
        self.title_label.setStyleSheet("color: #365ca3; background: transparent;")
        self.title_label.setWordWrap(True)
        self.title_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        header_layout.addWidget(self.title_label)

        main_layout.addWidget(self.header_widget)

        # ── Pestañas ────────────────────────────────────────────────
        tabs_container = QWidget()
        tabs_container.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        tabs_layout = QVBoxLayout(tabs_container)
        tabs_layout.setContentsMargins(8, 6, 8, 0)
        tabs_layout.setSpacing(0)

        self.tabs = QTabWidget()
        self.tab_individual = QWidget()
        self._setup_tab_individual()
        self.tabs.addTab(self.tab_individual, "👤 Individual")

        self.tab_masivo = QWidget()
        self._setup_tab_masivo()
        self.tabs.addTab(self.tab_masivo, "📊 Masivo (Excel)")

        tabs_layout.addWidget(self.tabs)
        main_layout.addWidget(tabs_container, 1)

        # ── Status bar ──────────────────────────────────────────────
        self.status_label = QLabel("Ingreso individual o carga Excel para flujo hasta formulario de boleta.")
        self.status_label.setAlignment(Qt.AlignCenter)
        self.status_label.setStyleSheet(
            "color: #365ca3; padding: 6px 10px; border-top: 1px solid #a3b4cb; background: #d4dce4;"
        )
        self.status_label.setWordWrap(True)
        main_layout.addWidget(self.status_label)

        self.set_custom_theme()
        self.setGeometry(100, 100, *self.VIEW_PROFILES["PC"][:2])
        # Aplicar perfil inicial tras el primer paint (splitter necesita tamaño real)
        QTimer.singleShot(0, lambda: self._apply_profile("PC"))

    def _on_view_changed(self, index):
        self.apply_view_mode(["PC", "Tablet", "Móvil"][index])

    def apply_view_mode(self, mode):
        self._apply_profile(mode)
        w, h = self.VIEW_PROFILES[mode][:2]
        self.resize(w, h)

    def _apply_profile(self, mode):
        if not hasattr(self, 'splitter') or not hasattr(self, 'log_text_masivo'):
            return
        profile = self.VIEW_PROFILES.get(mode, self.VIEW_PROFILES["PC"])
        w, h, font_sz, logo_h, splitter_code, splitter_sizes, log_h = profile
        splitter_orient = Qt.Horizontal if splitter_code == "H" else Qt.Vertical

        # Logo
        if self.logo_label and self._logo_path:
            pix = QPixmap(self._logo_path).scaledToHeight(logo_h, Qt.SmoothTransformation)
            if not pix.isNull():
                self.logo_label.setPixmap(pix)
            self.logo_label.setFixedHeight(logo_h + 2)

        # Fuente del título (escala proporcional y limpia)
        self.title_label.setFont(QFont("Segoe UI", font_sz, QFont.Bold))

        # Splitter
        self.splitter.setOrientation(splitter_orient)
        self.splitter.setSizes(splitter_sizes)

        # Log masivo
        self.log_text_masivo.setMaximumHeight(log_h)

    def _setup_tab_individual(self):
        layout = QVBoxLayout(self.tab_individual)
        self.splitter = QSplitter(Qt.Horizontal)

        # --- Panel izquierdo (con scroll para vistas pequeñas) ---
        left_inner = QWidget()
        left_layout = QVBoxLayout(left_inner)
        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setWidget(left_inner)
        left_scroll.setFrameShape(QScrollArea.NoFrame)
        left = left_scroll

        # Credenciales emisor
        cred = QGroupBox("🔑 Credenciales Emisor")
        cred_layout = QVBoxLayout()
        cred_layout.addWidget(QLabel("RUT emisor:*"))
        self.rut_input = QLineEdit()
        self.rut_input.setPlaceholderText("Ej: 76.123.456-7")
        cred_layout.addWidget(self.rut_input)
        cred_layout.addWidget(QLabel("Clave SII:*"))
        self.clave_input = QLineEdit()
        self.clave_input.setEchoMode(QLineEdit.Password)
        self.clave_input.setPlaceholderText("Clave SII")
        cred_layout.addWidget(self.clave_input)
        cred.setLayout(cred_layout)
        left_layout.addWidget(cred)

        # Datos destinatario
        dest = QGroupBox("👤 Destinatario")
        dest_layout = QVBoxLayout()
        dest_layout.addWidget(QLabel("RUT destinatario:*"))
        rut_row = QHBoxLayout()
        self.rut_dest_input = QLineEdit()
        self.rut_dest_input.setPlaceholderText("Ej: 12345678")
        self.dv_dest_input = QLineEdit()
        self.dv_dest_input.setPlaceholderText("DV")
        self.dv_dest_input.setMaximumWidth(45)
        rut_row.addWidget(self.rut_dest_input)
        rut_row.addWidget(QLabel("-"))
        rut_row.addWidget(self.dv_dest_input)
        dest_layout.addLayout(rut_row)
        dest_layout.addWidget(QLabel("Nombres:*"))
        self.nombres_dest_input = QLineEdit()
        self.nombres_dest_input.setPlaceholderText("Nombre / Razón Social")
        dest_layout.addWidget(self.nombres_dest_input)
        dest_layout.addWidget(QLabel("Domicilio:*"))
        self.domicilio_input = QLineEdit()
        self.domicilio_input.setPlaceholderText("Dirección")
        dest_layout.addWidget(self.domicilio_input)
        dest_layout.addWidget(QLabel("Región:"))
        self.region_input = QLineEdit()
        self.region_input.setPlaceholderText("Nombre o código de región")
        dest_layout.addWidget(self.region_input)
        dest_layout.addWidget(QLabel("Comuna:*"))
        self.comuna_input = QLineEdit()
        self.comuna_input.setPlaceholderText("Nombre o código de comuna")
        dest_layout.addWidget(self.comuna_input)
        dest.setLayout(dest_layout)
        left_layout.addWidget(dest)

        # Prestaciones
        prest = QGroupBox("💼 Prestaciones")
        prest_layout = QVBoxLayout()
        self.prestacion_inputs = []
        self.valor_inputs = []
        for i in range(1, 5):
            row = QHBoxLayout()
            p = QLineEdit()
            p.setPlaceholderText(f"Prestación {i}{'  *' if i == 1 else ''}")
            v = QLineEdit()
            v.setPlaceholderText(f"Valor {i}{'  *' if i == 1 else ''}")
            v.setMaximumWidth(100)
            row.addWidget(p)
            row.addWidget(v)
            prest_layout.addLayout(row)
            self.prestacion_inputs.append(p)
            self.valor_inputs.append(v)
        prest.setLayout(prest_layout)
        left_layout.addWidget(prest)

        # Opciones y botones
        opts = QGroupBox("⚙️ Opciones")
        opts_layout = QVBoxLayout()
        self.headless_check = QCheckBox("Modo sin interfaz (Headless)")
        opts_layout.addWidget(self.headless_check)
        opts.setLayout(opts_layout)
        left_layout.addWidget(opts)

        self.start_btn = QPushButton("🚀 Emitir Boleta")
        self.start_btn.clicked.connect(self.iniciar_proceso_individual)
        self.start_btn.setMinimumHeight(50)
        self.stop_btn = QPushButton("⏹️ Detener")
        self.stop_btn.clicked.connect(self.detener_proceso)
        self.stop_btn.setEnabled(False)
        left_layout.addWidget(self.start_btn)
        left_layout.addWidget(self.stop_btn)
        left_layout.addStretch()

        # --- Panel derecho ---
        right = QWidget()
        right_layout = QVBoxLayout(right)
        self.progress_bar = QProgressBar()
        right_layout.addWidget(self.progress_bar)
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        right_layout.addWidget(self.log_text)

        self.splitter.addWidget(left)
        self.splitter.addWidget(right)
        self.splitter.setSizes([380, 600])
        layout.addWidget(self.splitter)

    def _setup_tab_masivo(self):
        layout = QVBoxLayout(self.tab_masivo)

        top = QGroupBox("📁 Carga de Excel")
        top_layout = QVBoxLayout()
        top_layout.addWidget(QLabel(
            "Columnas requeridas: RUT | CLAVE | RUT-DEST | RUT-CV | NOMBRE-DEST | DOMICILIO-DEST | REGIÓN | COMUNA | PRESTACIÓN | VALOR"
        ))
        file_row = QHBoxLayout()
        self.excel_path_label = QLabel("No se ha seleccionado archivo")
        self.excel_path_label.setStyleSheet("color: #666; font-style: italic;")
        btn_load = QPushButton("📁 Cargar Excel")
        btn_load.clicked.connect(self._cargar_excel)
        file_row.addWidget(self.excel_path_label, 1)
        file_row.addWidget(btn_load)
        top_layout.addLayout(file_row)
        top.setLayout(top_layout)
        layout.addWidget(top)

        self.table_masivo = QTableWidget()
        self.table_masivo.setColumnCount(5)
        self.table_masivo.setHorizontalHeaderLabels(["RUT Emisor", "RUT Destinatario", "Prestación 1", "Monto", "Estado"])
        self.table_masivo.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.table_masivo)

        self.headless_masivo_check = QCheckBox("Modo sin interfaz (Headless)")
        self.headless_masivo_check.setChecked(True)
        layout.addWidget(self.headless_masivo_check)

        btn_row = QHBoxLayout()
        self.start_masivo_btn = QPushButton("🚀 Iniciar todo el Excel")
        self.start_masivo_btn.clicked.connect(self._iniciar_proceso_masivo)
        self.start_masivo_btn.setMinimumHeight(44)
        self.start_masivo_btn.setEnabled(False)
        self.stop_masivo_btn = QPushButton("⏹️ Detener")
        self.stop_masivo_btn.clicked.connect(self._detener_proceso_masivo)
        self.stop_masivo_btn.setEnabled(False)
        btn_row.addWidget(self.start_masivo_btn)
        btn_row.addWidget(self.stop_masivo_btn)
        layout.addLayout(btn_row)

        self.log_text_masivo = QTextEdit()
        self.log_text_masivo.setReadOnly(True)
        self.log_text_masivo.setMaximumHeight(180)
        layout.addWidget(self.log_text_masivo)

    def _cargar_excel(self):
        try:
            import pandas as pd
        except ImportError:
            QMessageBox.critical(self, "Error", "Instale pandas y openpyxl: pip install pandas openpyxl")
            return
        path, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar Excel", "", "Excel (*.xlsx *.xls);;CSV (*.csv)"
        )
        if not path:
            return
        self.excel_path = path
        self.excel_path_label.setText(os.path.basename(path))
        try:
            df = pd.read_csv(path) if path.lower().endswith(".csv") else pd.read_excel(path)
            # Normalizar nombres de columnas
            df.columns = [str(c).upper().strip().replace(" ", "_") for c in df.columns]

            def _col(nombres):
                for n in nombres:
                    if n in df.columns:
                        return n
                return None

            col_rut     = _col(["RUT", "RUT_EMISOR", "RUT_EMPRESA"])
            col_clave   = _col(["CLAVE", "CLAVE_SII", "PASSWORD"])
            col_rdest   = _col(["RUT_DEST", "RUT-DEST", "RUT_DESTINATARIO", "RUT_RECEPTOR"])
            col_dv      = _col(["RUT_CV", "RUT-CV", "DV_DEST", "DV", "DV_DESTINATARIO", "DV_RECEPTOR"])
            col_nombres = _col(["NOMBRE_DEST", "NOMBRE-DEST", "NOMBRES_DEST", "NOMBRE", "NOMBRES",
                                 "RAZON_SOCIAL", "NOMBRE_DESTINATARIO"])
            col_dom     = _col(["DOMICILIO_DEST", "DOMICILIO-DEST", "DOMICILIO", "DIRECCION", "DIR_DESTINATARIO"])
            col_region  = _col(["REGIÓN", "REGION", "COD_REGION"])
            col_comuna  = _col(["COMUNA", "COD_COMUNA"])

            # Prestación/Valor: soporta una columna única (solo para index 0) o múltiples numeradas
            cols_prest = []
            cols_valor = []
            for i in range(1, 5):
                if i == 1:
                    cp = _col(["PRESTACIÓN", "PRESTACION", "PRESTACION_1", "PRESTACION1", "GLOSA_1", "GLOSA1"])
                    cv = _col(["VALOR", "VALOR_1", "VALOR1", "MONTO_1", "MONTO1"])
                else:
                    cp = _col([f"PRESTACION_{i}", f"PRESTACION{i}", f"GLOSA_{i}", f"GLOSA{i}"])
                    cv = _col([f"VALOR_{i}", f"VALOR{i}", f"MONTO_{i}", f"MONTO{i}"])
                cols_prest.append(cp)
                cols_valor.append(cv)

            if not col_rut or not col_clave:
                QMessageBox.critical(self, "Error", "El Excel debe tener columnas RUT y CLAVE.")
                return
            if not col_rdest or not cols_prest[0]:
                QMessageBox.critical(self, "Error", "Faltan columnas: RUT_DEST y/o PRESTACION_1.")
                return

            self.excel_df = df
            self.excel_cols = {
                "rut": col_rut, "clave": col_clave,
                "rut_dest": col_rdest, "dv_dest": col_dv,
                "nombres_dest": col_nombres, "domicilio": col_dom,
                "region": col_region, "comuna": col_comuna,
                "prestaciones": cols_prest, "valores": cols_valor,
            }

            def _safe(val):
                """Convierte un valor de celda a str limpio; trata NaN/None como vacío."""
                if val is None:
                    return ""
                try:
                    import math
                    if isinstance(val, float) and math.isnan(val):
                        return ""
                except Exception:
                    pass
                return str(val).strip()

            self.table_masivo.setRowCount(len(df))
            for i, (_, row) in enumerate(df.iterrows()):
                self.table_masivo.setItem(i, 0, QTableWidgetItem(_safe(row[col_rut] if col_rut in row.index else "")))
                self.table_masivo.setItem(i, 1, QTableWidgetItem(_safe(row[col_rdest] if col_rdest in row.index else "")))
                self.table_masivo.setItem(i, 2, QTableWidgetItem(_safe(row[cols_prest[0]] if cols_prest[0] and cols_prest[0] in row.index else "")))
                self.table_masivo.setItem(i, 3, QTableWidgetItem(_safe(row[cols_valor[0]] if cols_valor[0] and cols_valor[0] in row.index else "")))
                self.table_masivo.setItem(i, 4, QTableWidgetItem("Pendiente"))
            self.start_masivo_btn.setEnabled(True)
            self._log_masivo(f"✅ Excel cargado: {len(df)} fila(s).")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo leer el archivo:\n{str(e)}")

    def _log_masivo(self, msg):
        self.log_text_masivo.append(msg)
        c = self.log_text_masivo.textCursor()
        c.movePosition(c.End)
        self.log_text_masivo.setTextCursor(c)

    def _iniciar_proceso_masivo(self):
        if self.table_masivo.rowCount() == 0:
            QMessageBox.warning(self, "Advertencia", "Cargue un Excel con al menos una fila.")
            return
        self.masivo_detenido = False
        self.masivo_resultados = []
        self.start_masivo_btn.setEnabled(False)
        self.stop_masivo_btn.setEnabled(True)
        self._log_masivo("🚀 Iniciando procesamiento masivo (secuencial)...")
        self._procesar_siguiente_fila_masivo(0)

    def _procesar_siguiente_fila_masivo(self, row_idx):
        if self.masivo_detenido:
            self._finalizar_masivo()
            return
        if row_idx >= self.table_masivo.rowCount():
            self._finalizar_masivo()
            return

        df = getattr(self, "excel_df", None)
        cols = getattr(self, "excel_cols", None)
        if df is None or cols is None:
            self._finalizar_masivo()
            return

        def _safe_cell(row, col):
            """Extrae celda de una fila pandas como str limpio; NaN/None → ''."""
            if not col or col not in row.index:
                return ""
            val = row[col]
            if val is None:
                return ""
            try:
                import math
                if isinstance(val, float) and math.isnan(val):
                    return ""
            except Exception:
                pass
            return str(val).strip()

        row = df.iloc[row_idx]
        rut = _safe_cell(row, cols["rut"])
        clave = _safe_cell(row, cols["clave"])

        if not rut or not clave:
            self.table_masivo.setItem(row_idx, 4, QTableWidgetItem("❌ Sin RUT/Clave"))
            self._log_masivo(f"Fila {row_idx + 1}: omitida (sin datos).")
            QTimer.singleShot(0, lambda: self._procesar_siguiente_fila_masivo(row_idx + 1))
            return

        # Armar boleta_data desde la fila
        prestaciones = []
        for i in range(4):
            glosa = _safe_cell(row, cols["prestaciones"][i])
            valor = _safe_cell(row, cols["valores"][i])
            if glosa:
                prestaciones.append({"glosa": glosa, "valor": valor})

        boleta_data = {
            "rut_dest":     _safe_cell(row, cols["rut_dest"]),
            "dv_dest":      _safe_cell(row, cols["dv_dest"]),
            "nombres_dest": _safe_cell(row, cols["nombres_dest"]),
            "domicilio":    _safe_cell(row, cols["domicilio"]),
            "region":       _safe_cell(row, cols["region"]),
            "comuna":       _safe_cell(row, cols["comuna"]),
            "prestaciones": prestaciones,
        }

        self.table_masivo.setItem(row_idx, 4, QTableWidgetItem("⏳ Procesando..."))
        self._log_masivo(f"▶️ Fila {row_idx + 1}/{self.table_masivo.rowCount()}: {rut}")
        headless = self.headless_masivo_check.isChecked()
        self.worker_masivo = self.automator.crear_worker_independiente(rut, clave, boleta_data, headless)
        self.worker_masivo.log_signal.connect(self._log_masivo)
        self.worker_masivo.finished_signal.connect(
            lambda exito, msj, r=row_idx: self._on_masivo_fila_finished(exito, msj, r)
        )
        self.worker_masivo.start()

    def _on_masivo_fila_finished(self, exito, mensaje, row_idx):
        estado = "✅ Éxito" if exito else "❌ Fallo"
        self.table_masivo.setItem(row_idx, 4, QTableWidgetItem(estado))
        self._log_masivo(f"🏁 Fila {row_idx + 1}: {mensaje}")
        self.masivo_resultados.append({
            "Fila":           row_idx + 1,
            "RUT Emisor":     self.table_masivo.item(row_idx, 0).text() if self.table_masivo.item(row_idx, 0) else "",
            "RUT Destinatario": self.table_masivo.item(row_idx, 1).text() if self.table_masivo.item(row_idx, 1) else "",
            "Prestación 1":   self.table_masivo.item(row_idx, 2).text() if self.table_masivo.item(row_idx, 2) else "",
            "Monto":          self.table_masivo.item(row_idx, 3).text() if self.table_masivo.item(row_idx, 3) else "",
            "Estado":         "Éxito" if exito else "Fallo",
            "Detalle":        mensaje,
        })
        self._procesar_siguiente_fila_masivo(row_idx + 1)

    def _finalizar_masivo(self):
        self._log_masivo("🎉 Fin del procesamiento masivo.")
        self.start_masivo_btn.setEnabled(True)
        self.stop_masivo_btn.setEnabled(False)
        self._exportar_reporte_masivo()

    def _exportar_reporte_masivo(self):
        if not self.masivo_resultados:
            return
        try:
            from datetime import datetime
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            nombre_base = os.path.splitext(os.path.basename(self.excel_path))[0] if self.excel_path else "reporte"
            directorio = os.path.dirname(self.excel_path) if self.excel_path else os.path.expanduser("~")
            nombre_archivo = f"{nombre_base}_reporte_{timestamp}.xlsx"
            ruta_reporte = os.path.join(directorio, nombre_archivo)

            df_reporte = pd.DataFrame(self.masivo_resultados)
            exitosos = (df_reporte["Estado"] == "Éxito").sum()
            fallidos  = (df_reporte["Estado"] == "Fallo").sum()

            with pd.ExcelWriter(ruta_reporte, engine="openpyxl") as writer:
                df_reporte.to_excel(writer, index=False, sheet_name="Resultados")
                ws = writer.sheets["Resultados"]

                # Ajustar ancho de columnas
                for col in ws.columns:
                    max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                    ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

                # Colorear filas por estado
                from openpyxl.styles import PatternFill, Font
                verde  = PatternFill("solid", fgColor="C6EFCE")
                rojo   = PatternFill("solid", fgColor="FFC7CE")
                bold   = Font(bold=True)
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    estado_cell = row[5]  # columna "Estado"
                    fill = verde if estado_cell.value == "Éxito" else rojo
                    for cell in row:
                        cell.fill = fill

                # Fila resumen al final
                ws.append([])
                ws.append(["", "", "", "", "Total procesados:", len(self.masivo_resultados)])
                ws.append(["", "", "", "", "Exitosos:",  int(exitosos)])
                ws.append(["", "", "", "", "Fallidos:",  int(fallidos)])
                for summary_row in ws.iter_rows(min_row=ws.max_row - 2, max_row=ws.max_row):
                    for cell in summary_row:
                        cell.font = bold

            self._log_masivo(f"📄 Reporte guardado: {nombre_archivo}")
            self._log_masivo(f"   ✅ Exitosos: {exitosos}  ❌ Fallidos: {fallidos}")
        except Exception as e:
            self._log_masivo(f"⚠️ No se pudo guardar el reporte: {e}")

    def _detener_proceso_masivo(self):
        self.masivo_detenido = True
        self._log_masivo("🛑 Deteniendo después del proceso actual...")
        if self.worker_masivo and self.worker_masivo.isRunning():
            self.automator.detener_worker_externo(self.worker_masivo)
        self.start_masivo_btn.setEnabled(True)
        self.stop_masivo_btn.setEnabled(False)

    def set_custom_theme(self):
        COLOR_FONDO = "#d4dce4"
        COLOR_LETRAS = "#365ca3"
        COLOR_DETALLES = "#a3b4cb"
        palette = QPalette()
        palette.setColor(QPalette.Window, QColor(COLOR_FONDO))
        palette.setColor(QPalette.WindowText, QColor(COLOR_LETRAS))
        palette.setColor(QPalette.Base, Qt.white)
        palette.setColor(QPalette.Text, QColor(COLOR_LETRAS))
        palette.setColor(QPalette.Button, QColor(COLOR_DETALLES))
        palette.setColor(QPalette.ButtonText, QColor(COLOR_LETRAS))
        palette.setColor(QPalette.Highlight, QColor(COLOR_LETRAS))
        palette.setColor(QPalette.HighlightedText, Qt.white)
        self.setPalette(palette)
        self.setStyleSheet(f"""
            QMainWindow {{ background-color: {COLOR_FONDO}; }}
            QGroupBox {{ border: 2px solid {COLOR_DETALLES}; border-radius: 8px; margin-top: 15px; font-weight: bold; color: {COLOR_LETRAS}; }}
            QGroupBox::title {{ subcontrol-origin: margin; left: 10px; padding: 0 5px; }}
            QLineEdit {{ padding: 8px; border: 1px solid {COLOR_DETALLES}; border-radius: 4px; background: white; color: {COLOR_LETRAS}; }}
            QPushButton {{ background: {COLOR_DETALLES}; color: {COLOR_LETRAS}; border: none; padding: 10px; border-radius: 5px; font-weight: bold; }}
            QPushButton:hover {{ background: #c0cedf; }}
            QPushButton:disabled {{ background: #ccc; color: #666; }}
            QTextEdit {{ background: white; color: {COLOR_LETRAS}; border: 1px solid {COLOR_DETALLES}; border-radius: 4px; }}
            QProgressBar {{ border: 1px solid {COLOR_DETALLES}; border-radius: 5px; text-align: center; background: white; }}
            QProgressBar::chunk {{ background: #4CAF50; }}
            QCheckBox {{ color: {COLOR_LETRAS}; }}
            QTabWidget::pane {{ border: 1px solid {COLOR_DETALLES}; background: {COLOR_FONDO}; }}
            QTabBar::tab {{ background: {COLOR_DETALLES}; color: {COLOR_LETRAS}; padding: 10px 20px; margin-right: 2px; }}
            QTabBar::tab:selected {{ background: white; font-weight: bold; }}
            QTableWidget {{ background: white; color: {COLOR_LETRAS}; }}
            QHeaderView::section {{ background: {COLOR_DETALLES}; padding: 6px; font-weight: bold; color: {COLOR_LETRAS}; }}
        """)

    def _actualizar_log(self, mensaje):
        self.log_text.append(mensaje)
        c = self.log_text.textCursor()
        c.movePosition(c.End)
        self.log_text.setTextCursor(c)

    def iniciar_proceso_individual(self):
        rut = self.rut_input.text().strip()
        clave = self.clave_input.text().strip()
        if not rut:
            QMessageBox.warning(self, "Advertencia", "Ingrese el RUT emisor.")
            return
        if not clave:
            QMessageBox.warning(self, "Advertencia", "Ingrese la clave SII.")
            return
        if not self.rut_dest_input.text().strip():
            QMessageBox.warning(self, "Advertencia", "Ingrese el RUT del destinatario.")
            return
        if not self.prestacion_inputs[0].text().strip():
            QMessageBox.warning(self, "Advertencia", "Ingrese al menos la Prestación 1.")
            return

        prestaciones = []
        for i in range(4):
            glosa = self.prestacion_inputs[i].text().strip()
            valor = self.valor_inputs[i].text().strip()
            if glosa:
                prestaciones.append({"glosa": glosa, "valor": valor})

        boleta_data = {
            "rut_dest":     self.rut_dest_input.text().strip(),
            "dv_dest":      self.dv_dest_input.text().strip(),
            "nombres_dest": self.nombres_dest_input.text().strip(),
            "domicilio":    self.domicilio_input.text().strip(),
            "region":       self.region_input.text().strip(),
            "comuna":       self.comuna_input.text().strip(),
            "prestaciones": prestaciones,
        }

        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.log_text.clear()
        self.progress_bar.setValue(0)
        self.status_label.setText("🚀 Iniciando proceso...")
        headless = self.headless_check.isChecked()
        self.worker_individual = self.automator.iniciar_proceso(rut, clave, boleta_data, headless)
        if self.worker_individual:
            self.worker_individual.log_signal.connect(self._actualizar_log)
            self.worker_individual.progress_signal.connect(self.progress_bar.setValue)
            self.worker_individual.finished_signal.connect(self.proceso_finalizado)
            self.worker_individual.start()
        else:
            QMessageBox.warning(self, "Advertencia", "Ya hay un proceso en ejecución.")
            self.start_btn.setEnabled(True)
            self.stop_btn.setEnabled(False)

    def detener_proceso(self):
        if self.worker_individual and self.worker_individual.isRunning():
            if self.automator.detener_proceso():
                self._actualizar_log("🛑 Proceso detenido.")
                self.status_label.setText("⏹️ Detenido.")
            self.start_btn.setEnabled(True)
            self.stop_btn.setEnabled(False)

    def proceso_finalizado(self, exito, mensaje):
        if exito:
            self._actualizar_log(f"✅ {mensaje}")
            self.status_label.setText("🎉 Proceso completado.")
        else:
            self._actualizar_log(f"❌ {mensaje}")
            self.status_label.setText("❌ Proceso falló.")
        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)

    def closeEvent(self, event):
        ind_activo = self.worker_individual and self.worker_individual.isRunning()
        mas_activo = self.worker_masivo and self.worker_masivo.isRunning()
        if ind_activo or mas_activo:
            reply = QMessageBox.question(
                self, "Confirmar salida", "Hay un proceso en ejecución. ¿Salir?",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
            )
            if reply == QMessageBox.Yes:
                if ind_activo:
                    self.detener_proceso()
                if mas_activo:
                    self._detener_proceso_masivo()
                event.accept()
            else:
                event.ignore()
        else:
            event.accept()


def main():
    try:
        from PyQt5.QtWidgets import QApplication as QtApp
        app = QtApp.instance()
        if app is None:
            app = QtApp(sys.argv)
        app.setStyle("Fusion")
        window = BoletaGUI()
        window.show()
        sys.exit(app.exec_())
    except Exception as e:
        err = f"Error al iniciar:\n{str(e)}\n\n"
        import traceback
        err += traceback.format_exc()
        try:
            from PyQt5.QtWidgets import QApplication as QtApp, QMessageBox
            (QtApp.instance() or QtApp(sys.argv))
            QMessageBox.critical(None, "Error", err)
        except Exception:
            with open(os.path.join(PROJECT_ROOT, "error_log.txt"), "w", encoding="utf-8") as f:
                f.write(err)
        sys.exit(1)


if __name__ == "__main__":
    main()
